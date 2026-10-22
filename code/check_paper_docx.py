"""论文草稿 DOCX 的结构与编号自检（文稿核验脚本，非题目求解代码）。"""
import os
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _project import find_project_root  # noqa: E402

from docx import Document  # noqa: E402

# paper_format 随 docx 工具链分发，未随本包提交：依次从包内 _docx_tools/ 与
# 环境变量 PAPER_TOOLS_DIR 查找，找不到时给出可操作提示而非堆栈报错。
_docx_tools = [Path(__file__).resolve().parent / "_docx_tools"]
_docx_tools += [Path(x) for x in os.environ.get("PAPER_TOOLS_DIR", "").split(os.pathsep) if x]
for _cand in _docx_tools:
    if (_cand / "paper_format.py").is_file():
        sys.path.insert(0, str(_cand))
        break
try:
    import paper_format as pf  # noqa: E402
except ImportError as exc:
    raise SystemExit(
        "未找到 paper_format 模块（DOCX 格式校验工具链）。\n"
        "请将 docx 工具链的 scripts 目录复制为本目录下的 _docx_tools/，"
        "或设置环境变量 PAPER_TOOLS_DIR 指向该目录。"
    ) from exc

p = find_project_root(__file__) / "论文草稿.docx"
if not p.is_file():
    raise SystemExit(f"未找到论文草稿：{p}")
doc = Document(str(p))
texts = [x.text.strip() for x in doc.paragraphs if x.text.strip()]
figs = len(doc._element.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"))
eqs = len(doc._element.findall(".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"))
print(f"段落(非空)={len(texts)}  图={figs}  表={len(doc.tables)}  公式={eqs}")

caps = [t for t in texts if re.match(r"^(图|表)\s*\d+", t)]
def check(kind):
    nums = [int(re.match(rf"^{kind}\s*(\d+)", t).group(1)) for t in caps if t.startswith(kind)]
    ok = nums == list(range(1, len(nums)+1))
    print(f"{kind}: {nums}  连续={ok}")
check("图"); check("表")

# 问题一小节标题
print("\n--- 3.x 小节 ---")
for t in texts:
    if re.match(r"^3(\.\d+)*\s", t):
        print("  ", t[:60])

print("\n--- 问题一图题注 ---")
for t in caps:
    if re.match(r"^图(6|7|8|9|10)\s", t): print("  ", t)

print("\n--- 问题一表题注 ---")
for t in caps:
    if re.match(r"^表(2|3)\s", t): print("  ", t)

# 结构核验（草稿无摘要/参考文献，仅看编号类问题）
errs = pf.validate_paper_structure(doc, "cumcm", require_rendered_pages=False)
print("\nvalidate_paper_structure 报错:")
for e in errs:
    print("  -", e)
