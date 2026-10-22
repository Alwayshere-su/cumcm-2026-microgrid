from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, linprog, milp
from scipy.sparse import lil_matrix, vstack


sys.path.insert(0, str(Path(__file__).resolve().parent))
from _project import find_project_root  # noqa: E402
ROOT = find_project_root(__file__)
DATA = ROOT / "data_processed" / "q1_typical_day.csv"
PARAMS = ROOT / "data_processed" / "q1_storage_parameters.json"
TEMPLATE = ROOT / "附件" / "附件5" / "result1.xlsx"
OUTPUT = ROOT / "result1.xlsx"
RESULTS = ROOT / "results"


def _indices(T: int) -> dict[str, np.ndarray]:
    start = 0
    out: dict[str, np.ndarray] = {}
    for name, size in (("buy", T), ("charge", T), ("discharge", T),
                       ("surplus", T), ("energy", T + 1), ("mode", T)):
        out[name] = np.arange(start, start + size)
        start += size
    return out


def _base_problem(df: pd.DataFrame, prm: dict) -> tuple:
    T = len(df)
    idx = _indices(T)
    n = int(idx["mode"][-1] + 1)
    max_energy = prm["max_charge_discharge_power_kw"] * prm["interval_hours"]
    eta_c = prm["charge_efficiency"]
    eta_d = prm["discharge_efficiency"]

    Aeq = lil_matrix((2 * T, n), dtype=float)
    beq = np.zeros(2 * T)
    net = df["load_kwh"].to_numpy() - df["pv_forecast_kwh"].to_numpy()
    for t in range(T):
        # buy + discharge - charge - surplus = load - pv
        Aeq[t, idx["buy"][t]] = 1
        Aeq[t, idx["discharge"][t]] = 1
        Aeq[t, idx["charge"][t]] = -1
        Aeq[t, idx["surplus"][t]] = -1
        beq[t] = net[t]
        # E(t+1) = E(t) + eta_c*charge - discharge/eta_d
        row = T + t
        Aeq[row, idx["energy"][t + 1]] = 1
        Aeq[row, idx["energy"][t]] = -1
        Aeq[row, idx["charge"][t]] = -eta_c
        Aeq[row, idx["discharge"][t]] = 1 / eta_d

    Aub = lil_matrix((2 * T, n), dtype=float)
    bub = np.zeros(2 * T)
    for t in range(T):
        # charge <= Bmax*y; discharge <= Bmax*(1-y)
        Aub[t, idx["charge"][t]] = 1
        Aub[t, idx["mode"][t]] = -max_energy
        Aub[T + t, idx["discharge"][t]] = 1
        Aub[T + t, idx["mode"][t]] = max_energy
        bub[T + t] = max_energy

    lower = np.zeros(n)
    upper = np.full(n, np.inf)
    upper[idx["charge"]] = max_energy
    upper[idx["discharge"]] = max_energy
    lower[idx["energy"]] = prm["operating_min_kwh"]
    upper[idx["energy"]] = prm["operating_max_kwh"]
    lower[idx["energy"][0]] = upper[idx["energy"][0]] = prm["initial_energy_kwh"]
    lower[idx["energy"][-1]] = upper[idx["energy"][-1]] = prm["terminal_energy_kwh"]
    upper[idx["mode"]] = 1
    integrality = np.zeros(n, dtype=int)
    integrality[idx["mode"]] = 1
    return idx, Aeq.tocsr(), beq, Aub.tocsr(), bub, Bounds(lower, upper), integrality


def _solve(c, Aeq, beq, Aub, bub, bounds, integrality):
    constraints = [LinearConstraint(Aeq, beq, beq),
                   LinearConstraint(Aub, -np.inf, bub)]
    result = milp(c=c, integrality=integrality, bounds=bounds,
                  constraints=constraints,
                  options={"presolve": True, "mip_rel_gap": 1e-9, "time_limit": 120})
    if not result.success:
        raise RuntimeError(f"MILP failed: status={result.status}, {result.message}")
    return result


def _lp_problem(df: pd.DataFrame, prm: dict) -> tuple:
    """Build the continuous Q1 model; variable order excludes the MILP mode flag."""
    T = len(df)
    idx = {}
    start = 0
    for name, size in (("buy", T), ("charge", T), ("discharge", T),
                       ("surplus", T), ("energy", T + 1)):
        idx[name] = np.arange(start, start + size)
        start += size
    n = start
    Aeq = lil_matrix((2 * T, n), dtype=float)
    beq = np.zeros(2 * T)
    net = df["load_kwh"].to_numpy() - df["pv_forecast_kwh"].to_numpy()
    eta_c = prm["charge_efficiency"]
    eta_d = prm["discharge_efficiency"]
    for t in range(T):
        Aeq[t, idx["buy"][t]] = 1.0
        Aeq[t, idx["discharge"][t]] = 1.0
        Aeq[t, idx["charge"][t]] = -1.0
        Aeq[t, idx["surplus"][t]] = -1.0
        beq[t] = net[t]
        row = T + t
        Aeq[row, idx["energy"][t + 1]] = 1.0
        Aeq[row, idx["energy"][t]] = -1.0
        Aeq[row, idx["charge"][t]] = -eta_c
        Aeq[row, idx["discharge"][t]] = 1.0 / eta_d

    max_interval_energy = prm["max_charge_discharge_power_kw"] * prm["interval_hours"]
    lower = np.zeros(n)
    upper = np.full(n, np.inf)
    upper[idx["charge"]] = max_interval_energy
    upper[idx["discharge"]] = max_interval_energy
    lower[idx["energy"]] = prm["operating_min_kwh"]
    upper[idx["energy"]] = prm["operating_max_kwh"]
    lower[idx["energy"][0]] = upper[idx["energy"][0]] = prm["initial_energy_kwh"]
    lower[idx["energy"][-1]] = upper[idx["energy"][-1]] = prm["terminal_energy_kwh"]
    return idx, Aeq.tocsr(), beq, lower, upper


def solve_q1() -> tuple[pd.DataFrame, dict]:
    df = pd.read_csv(DATA)
    prm = json.loads(PARAMS.read_text(encoding="utf-8"))
    if len(df) != 144:
        raise ValueError(f"Q1 must contain 144 intervals, got {len(df)}")
    idx, Aeq, beq, lower, upper = _lp_problem(df, prm)
    n = len(lower)
    price = df["price"].to_numpy()

    primary = np.zeros(n)
    primary[idx["buy"]] = price
    lp_bounds = list(zip(lower, upper))
    r1 = linprog(primary, A_eq=Aeq, b_eq=beq, bounds=lp_bounds, method="highs")
    if not r1.success:
        raise RuntimeError(f"Q1 primary LP failed: {r1.message}")
    cost_star = float(primary @ r1.x)
    cost_tol = 1e-8

    # 退化处理（条件触发）：命题1 保证互斥最优解存在，但求解器返回的退化解
    # 未必互斥。本问数据下首解已满足全天无同时充放电，故不启用第二级；
    # 仅当真出现同时充放（Σc_t·d_t > 0）时，才在锁定最优费用的前提下最小化吞吐。
    x = r1.x
    tie_break_used = float(np.sum(x[idx["charge"]] * x[idx["discharge"]])) > 1e-9
    if tie_break_used:
        cost_row = lil_matrix((1, n), dtype=float)
        cost_row[0, idx["buy"]] = price
        secondary = np.zeros(n)
        secondary[idx["charge"]] = 1.0
        secondary[idx["discharge"]] = 1.0
        r2 = linprog(secondary, A_ub=cost_row.tocsr(),
                     b_ub=[cost_star + cost_tol], A_eq=Aeq, b_eq=beq,
                     bounds=lp_bounds, method="highs")
        if not r2.success:
            raise RuntimeError(f"Q1 tie-break LP failed: {r2.message}")
        x = r2.x

    result = df.copy()
    result["buy_kwh"] = x[idx["buy"]]
    result["charge_kwh"] = x[idx["charge"]]
    result["discharge_kwh"] = x[idx["discharge"]]
    result["surplus_kwh"] = x[idx["surplus"]]
    result["energy_start_kwh"] = x[idx["energy"][:-1]]
    result["energy_end_kwh"] = x[idx["energy"][1:]]
    result["interval_cost_yuan"] = result["price"] * result["buy_kwh"]

    no_storage_buy = np.maximum(result["load_kwh"] - result["pv_forecast_kwh"], 0)
    no_storage_cost = float(np.sum(price * no_storage_buy))
    total_cost = float(result["interval_cost_yuan"].sum())
    balance_residual = (result["buy_kwh"] + result["pv_forecast_kwh"]
                        + result["discharge_kwh"] - result["load_kwh"]
                        - result["charge_kwh"] - result["surplus_kwh"])
    soc_residual = (result["energy_end_kwh"] - result["energy_start_kwh"]
                    - prm["charge_efficiency"] * result["charge_kwh"]
                    + result["discharge_kwh"] / prm["discharge_efficiency"])
    simultaneous = np.minimum(result["charge_kwh"], result["discharge_kwh"])
    product_certificate = float(np.sum(result["charge_kwh"] * result["discharge_kwh"]))

    # Same objective with binary charge/discharge flags: numerical tightness check only.
    m_idx, m_Aeq, m_beq, m_Aub, m_bub, m_bounds, integrality = _base_problem(df, prm)
    m_obj = np.zeros(len(m_bounds.lb))
    m_obj[m_idx["buy"]] = price
    milp_result = _solve(m_obj, m_Aeq, m_beq, m_Aub, m_bub, m_bounds, integrality)
    milp_cost = float(m_obj @ milp_result.x)

    # HiGHS bound marginals belong to the primary economic LP, not the cleanup LP.
    upper_marginal = np.asarray(r1.upper.marginals)
    lower_marginal = np.asarray(r1.lower.marginals)
    interior_energy = idx["energy"][1:-1]
    capacity_upper_value = float(max(0.0, -np.sum(upper_marginal[interior_energy])))
    soc_lower_value = float(max(0.0, np.sum(lower_marginal[interior_energy])))
    charge_limit_value = float(max(0.0, -np.sum(upper_marginal[idx["charge"]])))
    discharge_limit_value = float(max(0.0, -np.sum(upper_marginal[idx["discharge"]])))

    def perturbed_primary_cost(changes: dict[str, float]) -> float:
        changed = dict(prm)
        changed.update(changes)
        p_idx, p_Aeq, p_beq, p_lower, p_upper = _lp_problem(df, changed)
        p_obj = np.zeros(len(p_lower))
        p_obj[p_idx["buy"]] = price
        solved = linprog(p_obj, A_eq=p_Aeq, b_eq=p_beq,
                         bounds=list(zip(p_lower, p_upper)), method="highs")
        if not solved.success:
            raise RuntimeError(f"Q1 shadow-price perturbation failed: {solved.message}")
        return float(p_obj @ solved.x)

    delta = 1.0
    finite_difference = {
        "soc_upper_capacity": cost_star - perturbed_primary_cost({
            "operating_max_kwh": prm["operating_max_kwh"] + delta}),
        "soc_lower_relaxation": cost_star - perturbed_primary_cost({
            "operating_min_kwh": prm["operating_min_kwh"] - delta}),
        "charge_and_discharge_power": cost_star - perturbed_primary_cost({
            "max_charge_discharge_power_kw": prm["max_charge_discharge_power_kw"]
            + delta / prm["interval_hours"]}),
    }
    # The appendix specifies one common charge/discharge rating, so its finite
    # difference is compared with the sum of both sets of bound multipliers.
    dual_values = {
        "soc_upper_capacity": capacity_upper_value,
        "soc_lower_relaxation": soc_lower_value,
        "charge_and_discharge_power": charge_limit_value + discharge_limit_value,
    }
    max_fd_error = max(abs(finite_difference[k] - dual_values[k]) for k in dual_values)
    shadow_prices = {
        "soc_upper_capacity_value_yuan_per_kwh": capacity_upper_value,
        "soc_lower_relaxation_value_yuan_per_kwh": soc_lower_value,
        "charge_power_value_yuan_per_kwh": charge_limit_value,
        "discharge_power_value_yuan_per_kwh": discharge_limit_value,
        "common_power_rating_value_yuan_per_kw": (charge_limit_value + discharge_limit_value)
        * prm["interval_hours"],
        "active_soc_upper_intervals": int(np.sum(np.isclose(r1.x[interior_energy], upper[interior_energy], atol=1e-6))),
        "active_soc_lower_intervals": int(np.sum(np.isclose(r1.x[interior_energy], lower[interior_energy], atol=1e-6))),
        "active_charge_limit_intervals": int(np.sum(np.isclose(r1.x[idx["charge"]], upper[idx["charge"]], atol=1e-6))),
        "active_discharge_limit_intervals": int(np.sum(np.isclose(r1.x[idx["discharge"]], upper[idx["discharge"]], atol=1e-6))),
        "finite_difference_savings_yuan_for_one_unit": finite_difference,
        "max_dual_finite_difference_error_yuan": float(max_fd_error),
    }
    summary = {
        "model_type": "LP",
        "solver": "scipy.optimize.linprog (HiGHS)",
        "degeneracy_tie_break": {
            "used": bool(tie_break_used),
            "order": ["purchase_cost", "battery_throughput"]},
        "primary_status": r1.message,
        "tie_break_status": (r2.message if tie_break_used else "not_needed"),
        "primary_optimal_cost_yuan": cost_star,
        "primary_cost_tolerance_yuan": cost_tol,
        "tie_break_cost_gap_yuan": total_cost - cost_star,
        "battery_throughput_kwh": float(result["charge_kwh"].sum() + result["discharge_kwh"].sum()),
        "milp_crosscheck_cost_yuan": milp_cost,
        "milp_crosscheck_cost_gap_yuan": milp_cost - cost_star,
        "shadow_prices": shadow_prices,
        "total_purchase_kwh": float(result["buy_kwh"].sum()),
        "total_purchase_cost_yuan": total_cost,
        "no_storage_purchase_kwh": float(no_storage_buy.sum()),
        "no_storage_cost_yuan": no_storage_cost,
        "cost_saving_yuan": no_storage_cost - total_cost,
        "cost_saving_ratio": (no_storage_cost - total_cost) / no_storage_cost,
        "total_load_kwh": float(result["load_kwh"].sum()),
        "total_pv_kwh": float(result["pv_forecast_kwh"].sum()),
        "total_charge_kwh": float(result["charge_kwh"].sum()),
        "total_discharge_kwh": float(result["discharge_kwh"].sum()),
        "total_surplus_kwh": float(result["surplus_kwh"].sum()),
        "initial_energy_kwh": float(result["energy_start_kwh"].iloc[0]),
        "terminal_energy_kwh": float(result["energy_end_kwh"].iloc[-1]),
        "min_energy_kwh": float(min(result["energy_start_kwh"].min(), result["energy_end_kwh"].min())),
        "max_energy_kwh": float(max(result["energy_start_kwh"].max(), result["energy_end_kwh"].max())),
        "max_charge_kwh": float(result["charge_kwh"].max()),
        "max_discharge_kwh": float(result["discharge_kwh"].max()),
        "max_balance_residual_kwh": float(balance_residual.abs().max()),
        "max_soc_residual_kwh": float(soc_residual.abs().max()),
        "max_simultaneous_charge_discharge_kwh": float(simultaneous.max()),
        "sum_charge_times_discharge_kwh2": product_certificate,
        "active_charge_intervals": int((result["charge_kwh"] > 1e-6).sum()),
        "active_discharge_intervals": int((result["discharge_kwh"] > 1e-6).sum()),
        "active_surplus_intervals": int((result["surplus_kwh"] > 1e-6).sum()),
    }
    return result, summary


def write_outputs(result: pd.DataFrame, summary: dict) -> None:
    RESULTS.mkdir(exist_ok=True)
    result.to_csv(RESULTS / "q1_dispatch_10min.csv", index=False, encoding="utf-8-sig")
    (RESULTS / "q1_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    requested_starts = [10, 12, 14, 16, 18, 20]
    requested_slots = [hour * 6 + 1 for hour in requested_starts]
    table1 = result.set_index("slot").loc[requested_slots, ["buy_kwh"]].reset_index()
    table1.insert(0, "interval", [f"{h}:00-{h}:10" for h in requested_starts])
    table1 = pd.concat([table1, pd.DataFrame([
        {"interval": "全天购电量", "slot": np.nan, "buy_kwh": summary["total_purchase_kwh"]},
        {"interval": "全天购电费", "slot": np.nan, "buy_kwh": summary["total_purchase_cost_yuan"]},
    ])], ignore_index=True)
    table1.to_csv(RESULTS / "q1_table1_purchase.csv", index=False, encoding="utf-8-sig")

    bins = np.repeat(np.arange(6), 24)
    grouped = result.assign(bin=bins).groupby("bin", sort=True).agg(
        charge_kwh=("charge_kwh", "sum"), discharge_kwh=("discharge_kwh", "sum"))
    grouped.insert(0, "interval", ["0:00-4:00", "4:00-8:00", "8:00-12:00",
                                   "12:00-16:00", "16:00-20:00", "20:00-24:00"])
    grouped.to_csv(RESULTS / "q1_table2_storage.csv", index=False, encoding="utf-8-sig")

    shutil.copy2(TEMPLATE, OUTPUT)
    wb = load_workbook(OUTPUT)
    ws_buy = wb["计划购电量"]
    if ws_buy.max_row != 145 or ws_buy.max_column != 2:
        raise ValueError("result1 template purchase sheet structure changed")
    # 模板按区间起止排列，而附件1按区间终点记录：模板首行对应源slot=2，
    # 末行“0:00+1-0:10+1”在典型日重复假设下循环对应源slot=1。
    template_source_slots = list(range(2, 145)) + [1]
    mapping_rows = []
    for excel_row, source_slot in enumerate(template_source_slots, start=2):
        value = float(result.loc[result["slot"] == source_slot, "buy_kwh"].iloc[0])
        ws_buy.cell(excel_row, 2, round(value, 6))
        mapping_rows.append({
            "excel_row": excel_row,
            "template_interval": ws_buy.cell(excel_row, 1).value,
            "source_slot": source_slot,
            "source_endpoint": result.loc[result["slot"] == source_slot, "source_time"].iloc[0],
            "buy_kwh": value,
        })
    pd.DataFrame(mapping_rows).to_csv(
        RESULTS / "q1_result1_mapping_audit.csv", index=False, encoding="utf-8-sig")
    ws_storage = wb["充放电量"]
    for i, row in enumerate(grouped.itertuples(index=False), start=2):
        ws_storage.cell(i, 2, round(float(row.charge_kwh), 6))
        ws_storage.cell(i, 3, round(float(row.discharge_kwh), 6))
    ws_storage.cell(2, 5, round(summary["initial_energy_kwh"], 6))
    ws_storage.cell(3, 5, round(summary["terminal_energy_kwh"], 6))
    wb.save(OUTPUT)


def main() -> None:
    result, summary = solve_q1()
    write_outputs(result, summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
