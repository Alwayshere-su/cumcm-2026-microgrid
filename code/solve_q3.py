from __future__ import annotations

from typing import Optional
from pathlib import Path
import shutil

import numpy as np
import pandas as pd
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix
from openpyxl import load_workbook


def interpolate_anchor_curve(hourly_anchors: np.ndarray, issue_slot: int,
                             start_value: float) -> np.ndarray:
    """Interpolate 24 hourly point forecasts onto 10-minute interval endpoints."""
    anchors = np.asarray(hourly_anchors, float)
    if anchors.shape != (24,):
        raise ValueError("hourly_anchors must contain 24 values")
    anchor_slots = issue_slot + np.arange(1, 25) * 6
    target_slots = np.arange(issue_slot + 1, issue_slot + 24 * 6 + 1)
    xp = np.r_[issue_slot, anchor_slots]
    fp = np.r_[max(0.0, start_value), np.maximum(0.0, anchors)]
    return np.maximum(0.0, np.interp(target_slots, xp, fp))


def settlement_c(original: np.ndarray, adjusted: np.ndarray, price: np.ndarray,
                 emergency: np.ndarray) -> dict[str, float]:
    original = np.asarray(original, float)
    adjusted = np.asarray(adjusted, float)
    price = np.asarray(price, float)
    emergency = np.asarray(emergency, float)
    under = np.maximum(original - adjusted, 0.0)
    over = np.maximum(adjusted - original, 0.0)
    within = np.minimum(original, adjusted)
    within_cost = float(np.sum(price * within))
    under_cost = float(np.sum(0.5 * price * under))
    over_cost = float(np.sum(1.5 * price * over))
    emergency_cost = float(np.sum(5.0 * price * emergency))
    return {"within_cost": within_cost, "under_cost": under_cost,
            "over_cost": over_cost, "emergency_cost": emergency_cost,
            "under_kwh": float(under.sum()), "over_kwh": float(over.sum()),
            "total": within_cost + under_cost + over_cost + emergency_cost}


def solve_mpc(load: np.ndarray, pv: np.ndarray, price: np.ndarray, energy0: float,
              prm: dict, original_plan: Optional[np.ndarray] = None,
              terminal_target: Optional[float] = None,
              terminal_penalty: Optional[float] = None) -> dict[str, np.ndarray]:
    load, pv, price = map(lambda x: np.asarray(x, float), (load, pv, price))
    T = len(load)
    if not (len(pv) == len(price) == T):
        raise ValueError("load, pv and price lengths differ")
    use_dev = original_plan is not None
    dev_mask = np.isfinite(original_plan) if use_dev else np.zeros(T, dtype=bool)
    names = ["buy", "charge", "discharge", "surplus", "energy"]
    sizes = [T, T, T, T, T + 1]
    if use_dev:
        names += ["under", "over"]
        sizes += [T, T]
    if terminal_target is not None and terminal_penalty is not None:
        names += ["terminal_under", "terminal_over"]
        sizes += [1, 1]
    offsets, n = {}, 0
    for name, size in zip(names, sizes):
        offsets[name] = n
        n += size
    ix = lambda name, t=0: offsets[name] + t
    obj = np.zeros(n)
    obj[ix("buy"):ix("buy") + T] = price
    obj[ix("charge"):ix("surplus") + T] += 1e-9
    if use_dev:
        obj[ix("under"):ix("under") + T] = 0.5 * price * dev_mask
        obj[ix("over"):ix("over") + T] = 0.5 * price * dev_mask
    if "terminal_under" in offsets:
        obj[ix("terminal_under")] = terminal_penalty
        obj[ix("terminal_over")] = terminal_penalty

    rows = 2 * T + (T if use_dev else 0) + (1 if "terminal_under" in offsets else 0)
    A = lil_matrix((rows, n)); b = np.zeros(rows); r = 0
    for t in range(T):
        A[r, ix("buy", t)] = 1; A[r, ix("discharge", t)] = 1
        A[r, ix("charge", t)] = -1; A[r, ix("surplus", t)] = -1
        b[r] = load[t] - pv[t]; r += 1
        A[r, ix("energy", t + 1)] = 1; A[r, ix("energy", t)] = -1
        A[r, ix("charge", t)] = -prm["charge_efficiency"]
        A[r, ix("discharge", t)] = 1 / prm["discharge_efficiency"]
        r += 1
        if use_dev:
            if dev_mask[t]:
                A[r, ix("buy", t)] = 1; A[r, ix("under", t)] = 1
                A[r, ix("over", t)] = -1; b[r] = original_plan[t]
            r += 1
    if "terminal_under" in offsets:
        A[r, ix("energy", T)] = 1; A[r, ix("terminal_under")] = 1
        A[r, ix("terminal_over")] = -1; b[r] = terminal_target

    lo = np.zeros(n); hi = np.full(n, np.inf)
    block = prm["max_charge_discharge_power_kw"] * prm["interval_hours"]
    hi[ix("charge"):ix("charge") + T] = block
    hi[ix("discharge"):ix("discharge") + T] = block
    lo[ix("energy"):ix("energy") + T + 1] = prm["operating_min_kwh"]
    hi[ix("energy"):ix("energy") + T + 1] = prm["operating_max_kwh"]
    lo[ix("energy")] = hi[ix("energy")] = energy0
    if terminal_target is not None and terminal_penalty is None:
        lo[ix("energy", T)] = hi[ix("energy", T)] = terminal_target
    if use_dev:
        hi[ix("under"):ix("under") + T][~dev_mask] = 0.0
        hi[ix("over"):ix("over") + T][~dev_mask] = 0.0
    ans = milp(obj, bounds=Bounds(lo, hi), constraints=LinearConstraint(A.tocsr(), b, b),
               options={"presolve": True, "time_limit": 30})
    if not ans.success:
        raise RuntimeError(ans.message)
    out = {}
    for name, size in zip(names, sizes):
        out[name] = ans.x[offsets[name]:offsets[name] + size]
    return out


def execute_segment(plan: dict[str, np.ndarray], load: np.ndarray, pv: np.ndarray,
                    price: np.ndarray, energy0: float, prm: dict) -> pd.DataFrame:
    rows, energy = [], float(energy0)
    for t in range(len(load)):
        buy = max(0.0, float(plan["buy"][t]))
        charge = max(0.0, float(plan["charge"][t]))
        discharge = max(0.0, float(plan["discharge"][t]))
        net = buy + pv[t] + discharge - load[t] - charge
        emergency, surplus = max(-net, 0.0), max(net, 0.0)
        end = energy + prm["charge_efficiency"] * charge - discharge / prm["discharge_efficiency"]
        rows.append((t + 1, buy, charge, discharge, surplus, emergency,
                     energy, end, load[t], pv[t], price[t]))
        energy = end
    return pd.DataFrame(rows, columns=["slot", "buy_kwh", "charge_kwh", "discharge_kwh",
        "surplus_kwh", "emergency_kwh", "energy_start_kwh", "energy_end_kwh",
        "load_kwh", "pv_kwh", "price"])


def _day_values(frame: pd.DataFrame, day: pd.Timestamp, column: str) -> np.ndarray:
    x = frame[frame["date"] == day].sort_values("slot")
    if len(x) != 144:
        raise ValueError(f"missing 144 rows for {day.date()} in {column}")
    return x[column].to_numpy(float)


def _official_current_day(official: pd.DataFrame, actual: pd.DataFrame,
                          day: pd.Timestamp, issue_hour: int) -> np.ndarray:
    issue = day + pd.Timedelta(hours=issue_hour)
    f = official[official["issue_at"] == issue].sort_values("horizon_h")
    if len(f) != 24:
        raise ValueError(f"missing official forecast at {issue}")
    issue_slot = issue_hour * 6
    start = 0.0 if issue_slot == 0 else float(actual[actual.date == day].sort_values("slot").pv_kwh.iloc[issue_slot - 1] * 6)
    full = interpolate_anchor_curve(f.pv_forecast_kw.to_numpy(), issue_slot, start) / 6.0
    return full[:144 - issue_slot]


def causal_fused_pv(data: pd.DataFrame, forecasts: pd.DataFrame,
                    official: pd.DataFrame, day: pd.Timestamp, issue_hour: int,
                    window_days: int = 28) -> tuple[np.ndarray, dict]:
    """Bias-correct and blend forecasts using labels strictly before ``day``."""
    day = pd.Timestamp(day).normalize()
    start = issue_hour * 6
    raw = _official_current_day(official, data, day, issue_hour)
    current_lgb = _day_values(forecasts, day, "lgb_pv_kwh")[start:]
    prior_days = sorted(pd.Timestamp(x) for x in forecasts.loc[forecasts.date < day, "date"].unique())
    prior_days = prior_days[-window_days:]
    residuals, lgb_errors = [], []
    for past in prior_days:
        try:
            off = _official_current_day(official, data, past, issue_hour)
        except ValueError:
            continue
        actual = _day_values(data, past, "pv_kwh")[start:]
        lgb = _day_values(forecasts, past, "lgb_pv_kwh")[start:]
        residuals.append(off - actual)
        lgb_errors.append(lgb - actual)
    if not residuals:
        return current_lgb.copy(), {"official_weight": 0.0, "history_days": 0,
                                    "official_mae_kwh": np.nan, "lgb_mae_kwh": np.nan}
    residuals = np.vstack(residuals)
    bias = np.median(residuals, axis=0)
    corrected = np.maximum(0.0, raw - bias)
    official_mae = float(np.mean(np.abs(residuals - bias)))
    lgb_mae = float(np.mean(np.abs(np.vstack(lgb_errors))))
    eps = 1e-9
    weight = lgb_mae / (official_mae + lgb_mae + eps)
    fused = weight * corrected + (1.0 - weight) * current_lgb
    return np.maximum(0.0, fused), {"official_weight": float(weight),
        "history_days": len(residuals), "official_mae_kwh": official_mae,
        "lgb_mae_kwh": lgb_mae, "mean_bias_kwh": float(np.mean(bias))}


def run_q3_day(data: pd.DataFrame, forecasts: pd.DataFrame, official: pd.DataFrame,
               day: pd.Timestamp, energy0: float, prm: dict,
               updates: tuple[int, ...] = (6, 12, 18),
               pv_mode: str = "anchor") -> tuple[pd.DataFrame, dict]:
    """Run one day; each accepted MPC plan supplies both declaration and storage actions."""
    day = pd.Timestamp(day).normalize()
    actual = data[data.date == day].sort_values("slot")
    if len(actual) != 144:
        raise ValueError(f"missing actual day {day.date()}")
    load0 = _day_values(forecasts, day, "lgb_load_kwh")
    price0 = actual.price_fixed.to_numpy(float)
    if pv_mode in ("fused", "hybrid"):
        pv0, _ = causal_fused_pv(data, forecasts, official, day, 0)
    elif pv_mode == "anchor":
        pv0 = _official_current_day(official, actual, day, 0)
    else:
        raise ValueError("pv_mode must be 'anchor', 'fused' or 'hybrid'")
    next_day = day + pd.Timedelta(days=1)
    has_next = len(forecasts[forecasts.date == next_day]) == 144
    if has_next:
        load_next = _day_values(forecasts, next_day, "lgb_load_kwh")
        pv_next = _day_values(forecasts, next_day, "lgb_pv_kwh")
        price_next = price0.copy()
        load_h = np.r_[load0, load_next]; pv_h = np.r_[pv0, pv_next]; price_h = np.r_[price0, price_next]
    else:
        load_h, pv_h, price_h = load0, pv0, price0
    initial = solve_mpc(load_h, pv_h, price_h, energy0, prm, terminal_target=6000.0)
    original = initial["buy"][:144].copy()
    active = {k: initial[k][:144].copy() for k in ("buy", "charge", "discharge", "surplus")}
    boundaries = [0, 36, 72, 108, 144]
    executed, energy = [], float(energy0)
    for pos, start in enumerate(boundaries[:-1]):
        issue_hour = start // 6
        if start > 0 and issue_hour in updates:
            cur_load = load0[start:]
            use_fused = pv_mode == "fused" or (pv_mode == "hybrid" and issue_hour != 12)
            cur_pv = (causal_fused_pv(data, forecasts, official, day, issue_hour)[0]
                      if use_fused else
                      _official_current_day(official, actual, day, issue_hour))
            cur_price = price0[start:]
            if has_next:
                lh = np.r_[cur_load, load_next]; ph = np.r_[cur_pv, pv_next]; pr = np.r_[cur_price, price_next]
                ref = np.r_[original[start:], np.full(144, np.nan)]
            else:
                lh, ph, pr = cur_load, cur_pv, cur_price
                ref = original[start:]
            revised = solve_mpc(lh, ph, pr, energy, prm, original_plan=ref,
                                terminal_target=6000.0)
            ncur = 144 - start
            for k in active:
                active[k][start:] = revised[k][:ncur]
        end = boundaries[pos + 1]
        seg_plan = {k: active[k][start:end] for k in active}
        run = execute_segment(seg_plan, actual.load_kwh.to_numpy()[start:end],
                              actual.pv_kwh.to_numpy()[start:end], price0[start:end], energy, prm)
        run["slot"] = np.arange(start + 1, end + 1)
        executed.append(run); energy = float(run.energy_end_kwh.iloc[-1])
    out = pd.concat(executed, ignore_index=True)
    out.insert(0, "date", day)
    out["original_plan_kwh"] = original
    out["adjusted_buy_kwh"] = out["buy_kwh"]
    costs = settlement_c(original, out.adjusted_buy_kwh.to_numpy(), price0,
                         out.emergency_kwh.to_numpy())
    bal = (out.adjusted_buy_kwh + out.pv_kwh + out.discharge_kwh + out.emergency_kwh
           - out.load_kwh - out.charge_kwh - out.surplus_kwh)
    meta = {**costs, "terminal_soc_kwh": energy, "pv_mode": pv_mode,
            "max_balance_residual_kwh": float(np.abs(bal).max())}
    return out, meta


def run_q3_period(data: pd.DataFrame, forecasts: pd.DataFrame, official: pd.DataFrame,
                  dates, energy0: float, prm: dict,
                  updates: tuple[int, ...] = (6, 12, 18),
                  pv_mode: str = "anchor") -> tuple[pd.DataFrame, pd.DataFrame]:
    dispatch, metrics = [], []
    energy = float(energy0)
    for day in pd.DatetimeIndex(dates):
        run, meta = run_q3_day(data, forecasts, official, day, energy, prm, updates, pv_mode)
        dispatch.append(run)
        metrics.append({"date": day, **meta})
        energy = float(run.energy_end_kwh.iloc[-1])
    return pd.concat(dispatch, ignore_index=True), pd.DataFrame(metrics)


def _emergency_ranges(day: pd.DataFrame) -> list[tuple[str, float]]:
    active = day.emergency_kwh.to_numpy() > 1e-6
    values = day.emergency_kwh.to_numpy(); ans = []; i = 0
    def hm(k: int) -> str:
        minute = k * 10
        return f"{minute // 60}:{minute % 60:02d}" if minute < 1440 else f"0:{minute % 60:02d}+1"
    while i < len(day):
        if not active[i]:
            i += 1; continue
        j = i
        while j + 1 < len(day) and active[j + 1]:
            j += 1
        ans.append((f"{hm(i)}-{hm(j + 1)}", float(values[i:j + 1].sum())))
        i = j + 1
    return ans


def write_result3(out: pd.DataFrame, template: Path, target: Path) -> None:
    template, target = Path(template), Path(target)
    shutil.copy2(template, target)
    wb = load_workbook(target)
    days = sorted(pd.Timestamp(x) for x in out.date.unique())
    source_slots = list(range(2, 145)) + [1]
    for sheet, value_col in (("计划购电量", "original_plan_kwh"),
                             ("调整购电量", "adjusted_buy_kwh")):
        ws = wb[sheet]
        for row, day in enumerate(days, 2):
            x = out[out.date == day].set_index("slot")
            for col, slot in enumerate(source_slots, 2):
                ws.cell(row, col, round(float(x.loc[slot, value_col]), 6))
            ws.cell(row, 146, round(float(x[value_col].sum()), 6))
            if sheet == "计划购电量":
                fee = float((x.original_plan_kwh * x.price).sum())
            else:
                fee = settlement_c(x.original_plan_kwh.to_numpy(), x.adjusted_buy_kwh.to_numpy(),
                                   x.price.to_numpy(), x.emergency_kwh.to_numpy())["total"]
            ws.cell(row, 147, round(fee, 6))
    sw = wb["充放电量"]
    labels = [sw.cell(r, 2).value for r in range(2, 8)]
    if sw.max_row > 1:
        sw.delete_rows(2, sw.max_row - 1)
    row = 2
    for day in days:
        x = out[out.date == day].copy(); x["bin"] = (x.slot - 1) // 24
        grouped = x.groupby("bin").agg(charge=("charge_kwh", "sum"), discharge=("discharge_kwh", "sum"))
        for b in range(6):
            sw.cell(row, 1, day.to_pydatetime() if b == 0 else None)
            sw.cell(row, 2, labels[b])
            sw.cell(row, 3, round(float(grouped.loc[b, "charge"]), 6))
            sw.cell(row, 4, round(float(grouped.loc[b, "discharge"]), 6))
            if b == 0:
                sw.cell(row, 5, "0:00"); sw.cell(row, 6, round(float(x.energy_start_kwh.iloc[0]), 6))
            elif b == 1:
                sw.cell(row, 5, "24:00"); sw.cell(row, 6, round(float(x.energy_end_kwh.iloc[-1]), 6))
            row += 1
    ew = wb["紧急购电量"]
    if ew.max_row > 1:
        ew.delete_rows(2, ew.max_row - 1)
    row = 2
    for day in days:
        ranges = _emergency_ranges(out[out.date == day])
        for k, (interval, amount) in enumerate(ranges):
            ew.cell(row, 1, day.to_pydatetime() if k == 0 else None)
            ew.cell(row, 2, interval); ew.cell(row, 3, round(amount, 6)); row += 1
    wb.save(target)
