"""全面校验 result2.xlsx（问题二交付物）与方法三逐时段 dispatch 的一致性。

校验三张表：
  1) 计划购电量：334 天 × 144 时段逐格 + 全天购电量 + 全天购电费
  2) 充放电量 ：每天 6 个 4h 区的充/放电量 + 0:00 / 24:00 储电量
  3) 紧急购电量：全部紧急购电连续区间与电量
基准：results/q2_method3_dispatch.csv
"""

from __future__ import annotations

import datetime
import sys
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _project import find_project_root, add_script_path  # noqa: E402
ROOT = find_project_root(__file__)
add_script_path()
from solve_q2 import _ranges  # noqa: E402

DATES = list(pd.date_range("2025-02-01", "2025-12-31", freq="D"))
BLOCKS = [(1, 24, "0:00-4:00"), (25, 48, "4:00-8:00"), (49, 72, "8:00-12:00"),
          (73, 96, "12:00-16:00"), (97, 120, "16:00-20:00"), (121, 144, "20:00-24:00")]
TOL = 1e-6


def main() -> None:
    d = pd.read_csv(ROOT / "results" / "q2_method3_dispatch.csv", parse_dates=["date"])
    by_day = {day: g.sort_values("slot").reset_index(drop=True) for day, g in d.groupby("date")}
    wb = openpyxl.load_workbook(ROOT / "result2.xlsx", data_only=True)
    p_ws, s_ws, e_ws = wb[wb.sheetnames[0]], wb[wb.sheetnames[1]], wb[wb.sheetnames[2]]

    problems: list[str] = []

    # ---------- 1) 计划购电量 ----------
    assert p_ws.max_row == 1 + len(DATES), f"计划购电量行数 {p_ws.max_row}"
    assert p_ws.max_column == 147, f"计划购电量列数 {p_ws.max_column}"
    max_plan = max_all = max_fee = 0.0
    for i, day in enumerate(DATES):
        r = 2 + i
        cell_day = p_ws.cell(r, 1).value
        if not (isinstance(cell_day, datetime.datetime) and cell_day.date() == day.date()):
            problems.append(f"计划购电量 行{r} 日期={cell_day} 期望 {day.date()}")
        g = by_day[day]
        for slot in range(1, 145):
            col = 145 if slot == 1 else slot        # 模板列序：slot2..144 -> 列2..144, slot1 -> 列145
            xl = p_ws.cell(r, col).value
            exp = float(g.loc[slot - 1, "plan_buy_kwh"])
            max_plan = max(max_plan, abs(float(xl) - exp))
        max_all = max(max_all, abs(float(p_ws.cell(r, 146).value) - float(g.plan_buy_kwh.sum())))
        max_fee = max(max_fee, abs(float(p_ws.cell(r, 147).value) - float(g.plan_cost_yuan.sum())))

    # ---------- 2) 充放电量 ----------
    assert s_ws.max_row == 1 + len(DATES) * 6, f"充放电量行数 {s_ws.max_row}"
    max_chg = max_dis = max_soc0 = max_soc24 = 0.0
    for i, day in enumerate(DATES):
        g = by_day[day]
        blocks = [(float(g[g.slot.between(lo, hi)].charge_kwh.sum()),
                   float(g[g.slot.between(lo, hi)].discharge_kwh.sum())) for lo, hi, _ in BLOCKS]
        base = 2 + i * 6
        for b in range(6):
            r = base + b
            xc, xd = float(s_ws.cell(r, 3).value), float(s_ws.cell(r, 4).value)
            max_chg = max(max_chg, abs(xc - blocks[b][0]))
            max_dis = max(max_dis, abs(xd - blocks[b][1]))
        # 0:00 与 24:00 储电量写在每天第 1、2 行的第 6 列
        max_soc0 = max(max_soc0, abs(float(s_ws.cell(base, 6).value) - float(g.energy_start_kwh.iloc[0])))
        max_soc24 = max(max_soc24, abs(float(s_ws.cell(base + 1, 6).value) - float(g.energy_end_kwh.iloc[-1])))

    # ---------- 3) 紧急购电量 ----------
    rows = []
    for r in range(2, e_ws.max_row + 1):
        rows.append([e_ws.cell(r, 1).value, e_ws.cell(r, 2).value, float(e_ws.cell(r, 3).value)])
    exp_rows = []
    for day in DATES:
        for interval, amount in _ranges(by_day[day]):
            exp_rows.append([day, interval, float(amount)])
    max_em = 0.0
    if len(rows) != len(exp_rows):
        problems.append(f"紧急购电量记录数 {len(rows)} != 期望 {len(exp_rows)}")
    else:
        for (a_day, a_int, a_amt), (e_day, e_int, e_amt) in zip(rows, exp_rows):
            if a_int != e_int:
                problems.append(f"紧急区间不一致: xlsx {a_int} vs 期望 {e_int}")
            max_em = max(max_em, abs(a_amt - e_amt))
            if a_day is not None and not (isinstance(a_day, datetime.datetime) and a_day.date() == e_day.date()):
                problems.append(f"紧急日期不一致: {a_day} vs {e_day.date()}")

    total_plan = d.plan_buy_kwh.sum()
    total_em = d.emergency_kwh.sum()

    print("================ result2.xlsx 校验 ================")
    print(f"计划购电量：334 天 × 144 段，逐格最大偏差 = {max_plan:.3e}")
    print(f"           全天购电量最大偏差 = {max_all:.3e}；全天购电费最大偏差 = {max_fee:.3e}")
    print(f"充放电量  ：充/放最大偏差 = {max_chg:.3e} / {max_dis:.3e}")
    print(f"           0:00/24:00 储电量最大偏差 = {max_soc0:.3e} / {max_soc24:.3e}")
    print(f"紧急购电量：记录 {len(rows)} 条（期望 {len(exp_rows)}），电量最大偏差 = {max_em:.3e}")
    print(f"           全年紧急购电合计 = {total_em:,.3f} kWh（方法三 dispatch 同值）")
    print(f"           全年计划购电合计 = {total_plan:,.3f} kWh")
    print("---------------------------------------------------")
    if problems:
        print(f"发现问题 {len(problems)} 处：")
        for p in problems[:20]:
            print("  -", p)
    else:
        ok = (max_plan < TOL and max_all < TOL and max_fee < TOL and max_chg < TOL
              and max_dis < TOL and max_soc0 < TOL and max_soc24 < TOL and max_em < TOL)
        print("结论：三张表与方法三 dispatch 完全一致" if ok else "结论：数值存在超出容差的偏差")


if __name__ == "__main__":
    main()
