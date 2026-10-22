import json
import sys
import unittest
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parent))
from _project import find_project_root, add_source_path  # noqa: E402
ROOT = find_project_root(__file__)
add_source_path()
from solve_q1 import solve_q1, write_outputs

DATA = ROOT / "data_processed" / "q1_typical_day.csv"


@unittest.skipUnless(DATA.exists(), "原始附件未随仓库分发；运行 preprocess.py 后再测")
class Q1ModelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.result, cls.summary = solve_q1()
        cls.params = json.loads((ROOT / "data_processed" / "q1_storage_parameters.json").read_text(encoding="utf-8"))

    def test_energy_balance(self):
        r = self.result
        residual = r.buy_kwh + r.pv_forecast_kwh + r.discharge_kwh - r.load_kwh - r.charge_kwh - r.surplus_kwh
        self.assertLess(float(np.abs(residual).max()), 1e-5)

    def test_storage_state_and_bounds(self):
        r, p = self.result, self.params
        residual = r.energy_end_kwh - r.energy_start_kwh - p["charge_efficiency"] * r.charge_kwh + r.discharge_kwh / p["discharge_efficiency"]
        self.assertLess(float(np.abs(residual).max()), 1e-5)
        self.assertAlmostEqual(float(r.energy_start_kwh.iloc[0]), 6000, places=5)
        self.assertAlmostEqual(float(r.energy_end_kwh.iloc[-1]), 6000, places=5)
        self.assertGreaterEqual(float(r.energy_start_kwh.min()), 1200 - 1e-5)
        self.assertLessEqual(float(r.energy_end_kwh.max()), 10800 + 1e-5)

    def test_limits_and_exclusivity(self):
        limit = 5000 / 6
        self.assertLessEqual(float(self.result.charge_kwh.max()), limit + 1e-5)
        self.assertLessEqual(float(self.result.discharge_kwh.max()), limit + 1e-5)
        self.assertLess(float(np.minimum(self.result.charge_kwh, self.result.discharge_kwh).max()), 1e-6)

    def test_improves_no_storage_baseline(self):
        self.assertLessEqual(self.summary["total_purchase_cost_yuan"], self.summary["no_storage_cost_yuan"] + 1e-4)

    def test_canonical_solver_is_continuous_lp(self):
        """防止问题一主入口回退为含二元变量的 MILP。"""
        self.assertEqual(self.summary["model_type"], "LP")
        self.assertIn("linprog", self.summary["solver"])
        self.assertNotIn("mode_charge", self.result.columns)

    def test_degeneracy_tie_break_and_product_certificate(self):
        """主模型先锁定最低费用；仅当出现同时充放时才触发吞吐 tie-break。"""
        self.assertEqual(
            self.summary["degeneracy_tie_break"]["order"],
            ["purchase_cost", "battery_throughput"],
        )
        self.assertFalse(self.summary["degeneracy_tie_break"]["used"])
        self.assertLessEqual(
            self.summary["tie_break_cost_gap_yuan"],
            self.summary["primary_cost_tolerance_yuan"] + 1e-7,
        )
        self.assertLess(self.summary["sum_charge_times_discharge_kwh2"], 1e-6)

    def test_lp_milp_crosscheck_and_shadow_prices(self):
        """LP 紧性与关键资源边际价值必须进入可复核结果。"""
        self.assertLess(abs(self.summary["milp_crosscheck_cost_gap_yuan"]), 1e-5)
        shadow = self.summary["shadow_prices"]
        for key in (
            "soc_upper_capacity_value_yuan_per_kwh",
            "soc_lower_relaxation_value_yuan_per_kwh",
            "charge_power_value_yuan_per_kwh",
            "discharge_power_value_yuan_per_kwh",
        ):
            self.assertIn(key, shadow)
            self.assertGreaterEqual(shadow[key], -1e-9)
        self.assertLess(shadow["max_dual_finite_difference_error_yuan"], 1e-4)

    @unittest.skipUnless((ROOT / "附件" / "附件5").is_dir(), "结果模板未随仓库分发")
    def test_result1_template(self):
        write_outputs(self.result, self.summary)
        wb = load_workbook(ROOT / "result1.xlsx", data_only=False)
        self.assertEqual(wb.sheetnames, ["计划购电量", "充放电量"])
        self.assertEqual(wb["计划购电量"].max_row, 145)
        self.assertTrue(all(wb["计划购电量"].cell(i, 2).value is not None for i in range(2, 146)))
        self.assertEqual(wb["充放电量"]["E2"].value, 6000)
        self.assertEqual(wb["充放电量"]["E3"].value, 6000)
        # 模板10:00-10:10必须映射到附件1终点10:10（source slot 61）。
        ws = wb["计划购电量"]
        row = next(i for i in range(2, 146) if ws.cell(i, 1).value == "10:00-10:10")
        expected = float(self.result.loc[self.result.slot == 61, "buy_kwh"].iloc[0])
        self.assertAlmostEqual(float(ws.cell(row, 2).value), expected, places=5)


if __name__ == "__main__":
    unittest.main()
