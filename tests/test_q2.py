import json
import sys
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _project import find_project_root, add_source_path  # noqa: E402
ROOT = find_project_root(__file__)
add_source_path()
import solve_q2_method3 as method3

DATA = ROOT / "data_processed" / "microgrid_10min.csv"


@unittest.skipUnless(DATA.exists(), "原始附件未随仓库分发；运行四问求解后再测")
class Q2Method3Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.out = pd.read_csv(ROOT / "results" / "q2_method3_dispatch.csv", parse_dates=["date"])
        cls.summary = json.loads((ROOT / "results" / "q2_method3_summary.json").read_text(encoding="utf-8"))
        cls.actual = pd.read_csv(ROOT / "data_processed" / "microgrid_10min.csv", parse_dates=["date"])

    def test_official_model_contract(self):
        self.assertEqual(self.summary["model_id"], "q2_method3_blend_scenario_sp_dp")
        self.assertEqual(self.summary["forecast_model"], "blend")
        self.assertEqual(self.summary["retrain_days"], 7)
        self.assertEqual(self.summary["residual_window_days"], 30)
        self.assertEqual(self.summary["scenarios"], 30)
        self.assertAlmostEqual(self.summary["continuation_value_yuan_per_kwh"], 0.481548)

    def test_scope_and_causal_physics(self):
        self.assertEqual(len(self.out), 334 * 144)
        self.assertEqual(self.out.date.nunique(), 334)
        merged = self.out.merge(self.actual[["date", "slot", "load_kwh", "pv_kwh"]], on=["date", "slot"])
        bal = (merged.plan_buy_kwh + merged.pv_kwh + merged.discharge_kwh
               + merged.emergency_kwh - merged.load_kwh - merged.charge_kwh
               - merged.surplus_kwh)
        soc = (merged.energy_end_kwh - merged.energy_start_kwh
               - .9 * merged.charge_kwh + merged.discharge_kwh / .9)
        self.assertLess(float(np.abs(bal).max()), 1e-5)
        self.assertLess(float(np.abs(soc).max()), 1e-5)
        self.assertLess(float(np.minimum(merged.charge_kwh, merged.discharge_kwh).max()), 1e-8)

    def test_cross_day_soc_and_retention_bounds(self):
        starts = self.out.groupby("date").first().energy_start_kwh.to_numpy()[1:]
        ends = self.out.groupby("date").last().energy_end_kwh.to_numpy()[:-1]
        self.assertLess(float(np.abs(starts - ends).max()), 1e-5)
        self.assertGreaterEqual(float(self.out.retention_kwh.min()), 1200 - 1e-6)
        self.assertLessEqual(float(self.out.retention_kwh.max()), 10800 + 1e-6)

    @unittest.skipUnless((ROOT / "附件" / "附件5").is_dir(), "结果模板未随仓库分发")
    def test_method3_owns_official_outputs(self):
        self.assertTrue(callable(method3.save_method3_outputs))
        method3.save_method3_outputs({"dispatch": self.out, "summary": self.summary})
        canonical = json.loads((ROOT / "results" / "q2_summary.json").read_text(encoding="utf-8"))
        self.assertEqual(canonical["model_id"], "q2_method3_blend_scenario_sp_dp")
        wb = load_workbook(ROOT / "result2.xlsx", data_only=False)
        self.assertEqual(wb["计划购电量"].max_row, 335)
        self.assertTrue(all(wb["计划购电量"].cell(2, j).value is not None for j in range(2, 148)))
        self.assertEqual(wb["充放电量"].max_row, 1 + 334 * 6)

    def test_four_official_days_use_method3(self):
        report = json.loads((ROOT / "results" / "q2_four_days_final.json").read_text(encoding="utf-8"))
        self.assertEqual(set(report), {"2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"})
        for day in report.values():
            self.assertIn("table1_slots", day)
            self.assertIn("table2_storage", day)
            self.assertIn("table3_emergency", day)


if __name__ == "__main__":
    unittest.main()
