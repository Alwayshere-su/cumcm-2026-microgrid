import sys
import unittest
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _project import find_project_root, add_source_path  # noqa: E402
ROOT = find_project_root(__file__)
add_source_path()

from solve_q3 import (interpolate_anchor_curve, settlement_c, solve_mpc,
                      execute_segment, run_q3_day, run_q3_period,
                      causal_fused_pv, write_result3)


PRM = {
    "max_charge_discharge_power_kw": 6,
    "interval_hours": 1 / 6,
    "charge_efficiency": 0.9,
    "discharge_efficiency": 0.9,
    "operating_min_kwh": 0,
    "operating_max_kwh": 10,
}


class Q3Tests(unittest.TestCase):
    def test_anchor_interpolation_hits_every_hourly_anchor(self):
        anchors = np.arange(1.0, 25.0)
        curve = interpolate_anchor_curve(anchors, issue_slot=0, start_value=0.0)
        self.assertEqual(len(curve), 144)
        np.testing.assert_allclose(curve[np.arange(6, 145, 6) - 1], anchors)

    def test_settlement_c_prices_under_and_over_purchase_once(self):
        plan = np.array([10.0, 10.0])
        adjusted = np.array([8.0, 13.0])
        cost = settlement_c(plan, adjusted, np.ones(2), np.zeros(2))
        # under: 8 + .5*(2) = 9; over: 10 + 1.5*(3) = 14.5
        self.assertAlmostEqual(cost["total"], 23.5)
        self.assertAlmostEqual(cost["under_kwh"], 2.0)
        self.assertAlmostEqual(cost["over_kwh"], 3.0)

    def test_mpc_obeys_balance_soc_and_terminal_target(self):
        load = np.array([2.0, 2.0, 2.0, 2.0])
        pv = np.zeros(4)
        sol = solve_mpc(load, pv, np.ones(4), 5.0, PRM,
                        original_plan=np.full(4, 2.0), terminal_target=5.0)
        residual = sol["buy"] + sol["discharge"] - sol["charge"] - sol["surplus"] - load
        self.assertLess(np.abs(residual).max(), 1e-7)
        self.assertAlmostEqual(sol["energy"][0], 5.0, places=7)
        self.assertAlmostEqual(sol["energy"][-1], 5.0, places=7)

    def test_execution_keeps_declared_storage_actions_and_closes_error(self):
        plan = {"buy": np.array([1.0]), "charge": np.array([0.4]),
                "discharge": np.array([0.0]), "surplus": np.array([0.0])}
        run = execute_segment(plan, np.array([2.0]), np.array([0.0]),
                              np.array([1.0]), 5.0, PRM)
        self.assertAlmostEqual(run.loc[0, "charge_kwh"], 0.4)
        self.assertAlmostEqual(run.loc[0, "emergency_kwh"], 1.4)
        self.assertAlmostEqual(run.loc[0, "energy_end_kwh"], 5.36)
        balance = (run.buy_kwh + run.pv_kwh + run.discharge_kwh + run.emergency_kwh
                   - run.load_kwh - run.charge_kwh - run.surplus_kwh)
        self.assertLess(np.abs(balance).max(), 1e-9)

    def test_nan_reference_marks_terminal_buffer_without_adjustment_fee(self):
        load = np.ones(3)
        sol = solve_mpc(load, np.zeros(3), np.array([1.0, 1.0, 0.1]), 5.0, PRM,
                        original_plan=np.array([1.0, 1.0, np.nan]), terminal_target=5.0)
        self.assertTrue(np.isfinite(sol["buy"]).all())
        self.assertEqual(len(sol["under"]), 3)
        self.assertAlmostEqual(sol["under"][2] + sol["over"][2], 0.0, places=7)

    @unittest.skipUnless((ROOT / "data_processed" / "microgrid_10min.csv").exists(),
                         "原始附件未随仓库分发")
    def test_real_day_runner_returns_one_causal_day(self):
        import pandas as pd
        data = pd.read_csv(ROOT / "data_processed/microgrid_10min.csv", parse_dates=["date", "timestamp_end"])
        fc = pd.read_csv(ROOT / "data_processed/q2_forecasts.csv", parse_dates=["date"])
        official = pd.read_csv(ROOT / "data_processed/pv_forecast_hourly.csv",
                               parse_dates=["issue_at", "target_end"])
        import json
        prm = json.loads((ROOT / "data_processed/q1_storage_parameters.json").read_text(encoding="utf-8"))
        out, meta = run_q3_day(data, fc, official, pd.Timestamp("2025-03-20"), 6000.0,
                               prm, updates=(6, 12, 18), pv_mode="fused")
        self.assertEqual(len(out), 144)
        self.assertTrue((out.adjusted_buy_kwh >= -1e-8).all())
        self.assertLess(meta["max_balance_residual_kwh"], 1e-7)
        self.assertEqual(meta["pv_mode"], "fused")

        _, hybrid_meta = run_q3_day(data, fc, official, pd.Timestamp("2025-03-20"),
                                    6000.0, prm, updates=(6, 12), pv_mode="hybrid")
        self.assertEqual(hybrid_meta["pv_mode"], "hybrid")

    @unittest.skipUnless((ROOT / "data_processed" / "microgrid_10min.csv").exists(),
                         "原始附件未随仓库分发")
    def test_period_runner_carries_soc_between_days(self):
        import pandas as pd
        data = pd.read_csv(ROOT / "data_processed/microgrid_10min.csv", parse_dates=["date", "timestamp_end"])
        fc = pd.read_csv(ROOT / "data_processed/q2_forecasts.csv", parse_dates=["date"])
        official = pd.read_csv(ROOT / "data_processed/pv_forecast_hourly.csv",
                               parse_dates=["issue_at", "target_end"])
        import json
        prm = json.loads((ROOT / "data_processed/q1_storage_parameters.json").read_text(encoding="utf-8"))
        out, daily = run_q3_period(data, fc, official,
                                   pd.date_range("2025-03-20", "2025-03-21"),
                                   6000.0, prm, updates=(6, 12))
        self.assertEqual(len(out), 288)
        self.assertEqual(len(daily), 2)
        self.assertAlmostEqual(out.iloc[143].energy_end_kwh,
                               out.iloc[144].energy_start_kwh, places=7)

    @unittest.skipUnless((ROOT / "data_processed" / "microgrid_10min.csv").exists(),
                         "原始附件未随仓库分发")
    def test_fused_forecast_does_not_read_current_future_actuals(self):
        import pandas as pd
        data = pd.read_csv(ROOT / "data_processed/microgrid_10min.csv", parse_dates=["date", "timestamp_end"])
        fc = pd.read_csv(ROOT / "data_processed/q2_forecasts.csv", parse_dates=["date"])
        official = pd.read_csv(ROOT / "data_processed/pv_forecast_hourly.csv",
                               parse_dates=["issue_at", "target_end"])
        day = pd.Timestamp("2025-03-20")
        a, _ = causal_fused_pv(data, fc, official, day, 6, window_days=28)
        changed = data.copy()
        mask = (changed.date == day) & (changed.slot > 36)
        changed.loc[mask, "pv_kwh"] = changed.loc[mask, "pv_kwh"] + 99999
        b, _ = causal_fused_pv(changed, fc, official, day, 6, window_days=28)
        np.testing.assert_allclose(a, b)

    @unittest.skipUnless((ROOT / "附件" / "附件5").is_dir(), "结果模板未随仓库分发")
    def test_result3_writer_preserves_four_sheet_contract(self):
        import pandas as pd
        import tempfile
        from openpyxl import load_workbook
        out = pd.read_csv(ROOT / "results/q3_final_dispatch.csv", parse_dates=["date"])
        with tempfile.TemporaryDirectory() as td:
            target = Path(td) / "result3.xlsx"
            write_result3(out, ROOT / "附件/附件5/result3.xlsx", target)
            wb = load_workbook(target, data_only=True)
            self.assertEqual(wb.sheetnames, ["计划购电量", "调整购电量", "充放电量", "紧急购电量"])
            self.assertEqual(wb["计划购电量"].max_row, 335)
            self.assertEqual(wb["调整购电量"].max_column, 147)
            self.assertEqual(wb["充放电量"].max_row, 334 * 6 + 1)


if __name__ == "__main__":
    unittest.main()
